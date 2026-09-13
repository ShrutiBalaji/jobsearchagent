"""Append qualifying jobs to a persistent Excel tracker."""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from job_agent.config import load_config
from job_agent.models import JobRecord
from job_agent.scoring import format_experience_needed, format_skills_missing
from job_agent.utils.salary import ParsedSalary, format_compensation

logger = logging.getLogger("job_agent.excel")

COLUMNS = [
    "Priority",
    "Company name",
    "URL to apply",
    "Role name",
    "Salary",
    "Match score",
    "Experience needed",
    "Location",
    "Date posted",
    "Skill missing",
    "Applied",
]

URL_COLUMN = 3
HEADER_FONT = Font(bold=True)


def _salary_text(job: JobRecord) -> str:
    if job.compensation_unknown and job.salary_min is None:
        return "Not listed"
    parsed = ParsedSalary(
        salary_min=job.salary_min,
        salary_max=job.salary_max,
        salary_currency=job.salary_currency,
        salary_period=job.salary_period or "yearly",
        compensation_unknown=job.compensation_unknown,
        hourly_estimated_annual=job.salary_period == "hourly",
    )
    return format_compensation(parsed)


def _location_text(job: JobRecord) -> str:
    if job.location:
        return job.location
    parts = [job.city, job.state, job.country]
    return ", ".join(p for p in parts if p) or "Unknown"


def _migrate_header(ws: Worksheet) -> None:
    """Upgrade an older Excel header to the current column layout."""
    header = [cell.value for cell in ws[1]]
    if header == COLUMNS:
        return

    # Previous layout without "Date posted"
    legacy = [
        "Company name",
        "URL to apply",
        "Role name",
        "Salary",
        "Match score",
        "Location",
        "Applied",
    ]
    if header == legacy:
        ws.insert_cols(7)
        ws.cell(row=1, column=7, value="Date posted")
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=7, value="")
        _style_header(ws)
        header = [cell.value for cell in ws[1]]

    # Layout with "Date posted" but without "Experience needed"
    without_experience = [
        "Company name",
        "URL to apply",
        "Role name",
        "Salary",
        "Match score",
        "Location",
        "Date posted",
        "Applied",
    ]
    if header == without_experience:
        ws.insert_cols(6)
        ws.cell(row=1, column=6, value="Experience needed")
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=6, value="")
        _style_header(ws)
        header = [cell.value for cell in ws[1]]

    without_skills = [
        "Company name",
        "URL to apply",
        "Role name",
        "Salary",
        "Match score",
        "Experience needed",
        "Location",
        "Date posted",
        "Applied",
    ]
    if header == without_skills:
        ws.insert_cols(9)
        ws.cell(row=1, column=9, value="Skill missing")
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=9, value="")
        _style_header(ws)
        header = [cell.value for cell in ws[1]]

    without_priority = [
        "Company name",
        "URL to apply",
        "Role name",
        "Salary",
        "Match score",
        "Experience needed",
        "Location",
        "Date posted",
        "Skill missing",
        "Applied",
    ]
    if header == without_priority:
        ws.insert_cols(1)
        ws.cell(row=1, column=1, value="Priority")
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=1, value="")
        _style_header(ws)
        return

    ws.delete_rows(1)
    ws.insert_rows(1)
    for col, name in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=col, value=name)
    _style_header(ws)


def _ensure_workbook(path: Path) -> tuple[Workbook, Worksheet, bool]:
    """Load existing workbook or create a new one. Returns (wb, ws, created)."""
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
        if ws.max_row == 0:
            ws.append(COLUMNS)
            _style_header(ws)
            return wb, ws, False
        _migrate_header(ws)
        return wb, ws, False

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "AI Job Matches"
    ws.append(COLUMNS)
    _style_header(ws)
    return wb, ws, True


def _style_header(ws: Worksheet) -> None:
    for cell in ws[1]:
        cell.font = HEADER_FONT


def _existing_application_urls(ws: Worksheet) -> set[str]:
    urls: set[str] = set()
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row=row, column=URL_COLUMN).value
        if value:
            urls.add(str(value).strip())
    return urls


def _date_posted_text(job: JobRecord) -> str:
    if job.date_posted:
        return job.date_posted.isoformat()
    return "Unknown"


def _experience_needed_text(job: JobRecord) -> str:
    combined = f"{job.title} {job.description} {job.requirements}"
    return format_experience_needed(combined)


def _skills_missing_text(job: JobRecord) -> str:
    if job.missing_requirements and job.missing_requirements != "None identified":
        return job.missing_requirements
    combined = f"{job.title} {job.description} {job.requirements}"
    return format_skills_missing(combined, load_config())


def _priority_text(job: JobRecord) -> str:
    if job.search_priority_label:
        return job.search_priority_label
    if job.search_priority_tier:
        return f"Tier {job.search_priority_tier}"
    return ""


def _job_row(job: JobRecord) -> list:
    return [
        _priority_text(job),
        job.company,
        job.application_url or job.url,
        job.title,
        _salary_text(job),
        job.match_score,
        _experience_needed_text(job),
        _location_text(job),
        _date_posted_text(job),
        _skills_missing_text(job),
        "",
    ]


def append_jobs_to_excel(path: str | Path, jobs: list[JobRecord]) -> tuple[int, Path]:
    """Sync qualifying jobs to the Excel tracker.

    Writes all provided jobs that are not already in the file (matched by apply URL).
    When the file is missing, creates it and writes every qualifying job passed in.
    Returns (rows_written, absolute_path).
    """
    resolved = Path(path).resolve()
    if not jobs:
        logger.info("No qualifying jobs to write to Excel")
        return 0, resolved

    wb, ws, created = _ensure_workbook(resolved)
    existing_urls = _existing_application_urls(ws)

    written = 0
    for job in jobs:
        apply_url = (job.application_url or job.url or "").strip()
        if not apply_url or apply_url in existing_urls:
            continue
        ws.append(_job_row(job))
        existing_urls.add(apply_url)
        written += 1

    if written or created:
        wb.save(resolved)
        if created:
            logger.info("Created Excel tracker with %d job(s) at %s", written, resolved)
        else:
            logger.info("Wrote %d new job(s) to %s", written, resolved)
    else:
        logger.info("No new rows written — all jobs already in Excel")

    return written, resolved


def rebuild_jobs_in_excel(path: str | Path, jobs: list[JobRecord]) -> tuple[int, Path]:
    """Replace the Excel tracker with the current qualifying job list."""
    resolved = Path(path).resolve()
    resolved.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "AI Job Matches"
    ws.append(COLUMNS)
    _style_header(ws)

    seen_urls: set[str] = set()
    written = 0
    for job in jobs:
        apply_url = (job.application_url or job.url or "").strip()
        if not apply_url or apply_url in seen_urls:
            continue
        ws.append(_job_row(job))
        seen_urls.add(apply_url)
        written += 1

    wb.save(resolved)
    logger.info("Rebuilt Excel tracker with %d job(s) at %s", written, resolved)
    return written, resolved
