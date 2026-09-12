"""Tests for Excel export."""

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from job_agent.excel_export import COLUMNS, append_jobs_to_excel
from job_agent.models import JobRecord, RemoteType, ScoreBreakdown


def _sample_job(**kwargs) -> JobRecord:
    defaults = dict(
        company="Anthropic",
        title="Machine Learning Engineer",
        url="https://example.com/jobs/1",
        application_url="https://example.com/jobs/1/apply",
        source="test",
        match_score=85.5,
        location="New York, NY",
        salary_min=150_000,
        salary_max=200_000,
        compensation_unknown=False,
        score_breakdown=ScoreBreakdown(),
        remote_type=RemoteType.REMOTE,
        date_posted=date.today(),
    )
    defaults.update(kwargs)
    return JobRecord(**defaults)


def test_create_excel_with_jobs(tmp_path: Path):
    path = tmp_path / "jobs.xlsx"
    count, saved = append_jobs_to_excel(path, [_sample_job()])
    assert count == 1
    assert saved.exists()

    wb = load_workbook(saved)
    ws = wb.active
    assert [cell.value for cell in ws[1]] == COLUMNS
    assert ws.cell(row=2, column=1).value == "Anthropic"
    assert ws.cell(row=2, column=7).value == "New York, NY"
    assert ws.cell(row=2, column=8).value == date.today().isoformat()
    assert ws.cell(row=2, column=10).value in ("", None)


def test_append_skips_duplicate_urls(tmp_path: Path):
    path = tmp_path / "jobs.xlsx"
    job = _sample_job()
    append_jobs_to_excel(path, [job])

    duplicate = _sample_job(title="Different Title")
    count, _ = append_jobs_to_excel(path, [duplicate])
    assert count == 0

    wb = load_workbook(path)
    assert wb.active.max_row == 2


def test_experience_needed_column(tmp_path: Path):
    path = tmp_path / "jobs.xlsx"
    job = _sample_job(
        description="5+ years of machine learning experience required. Python, PyTorch.",
    )
    append_jobs_to_excel(path, [job])

    wb = load_workbook(path)
    assert wb.active.cell(row=2, column=6).value == "5+ years"


def test_rebuild_excel_from_all_qualifying_jobs(tmp_path: Path):
    path = tmp_path / "jobs.xlsx"
    jobs = [
        _sample_job(
            company="Anthropic",
            application_url="https://example.com/jobs/1/apply",
            url="https://example.com/jobs/1",
        ),
        _sample_job(
            company="OpenAI",
            title="AI Engineer",
            application_url="https://example.com/jobs/2/apply",
            url="https://example.com/jobs/2",
        ),
    ]
    count, saved = append_jobs_to_excel(path, jobs)
    assert count == 2
    assert saved.exists()

    saved.unlink()
    assert not path.exists()

    count, saved = append_jobs_to_excel(path, jobs)
    assert count == 2
    assert saved.exists()

    wb = load_workbook(saved)
    assert wb.active.max_row == 3


def test_skills_missing_column(tmp_path: Path):
    path = tmp_path / "jobs.xlsx"
    job = _sample_job(
        description="Requires Kubernetes, Java, and PyTorch in production ML systems.",
    )
    append_jobs_to_excel(path, [job])

    wb = load_workbook(path)
    skills = wb.active.cell(row=2, column=9).value
    assert skills
    assert "Kubernetes" in skills or "Java" in skills


def test_unknown_salary_shown(tmp_path: Path):
    path = tmp_path / "jobs.xlsx"
    job = _sample_job(salary_min=None, salary_max=None, compensation_unknown=True)
    append_jobs_to_excel(path, [job])

    wb = load_workbook(path)
    assert wb.active.cell(row=2, column=4).value == "Not listed"
